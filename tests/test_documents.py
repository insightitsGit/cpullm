"""A file or a hand-built document becomes the same paragraph chunks."""

import csv
import json
import threading
import zipfile
from pathlib import Path

import docx
import openpyxl
import pytest
from pypdf import PdfReader

from cpu_llm.documents import chunks_for, claim_source_names, open_source, paragraphs_from_text, structured_paragraphs
from cpu_llm.documents import SourceDocument, _slug
from cpu_llm.relations import add_edges, load_edges


def test_a_developer_names_the_paragraphs_directly():
    document = SourceDocument(
        "field-guide",
        (
            "The European lobster is Homarus gammarus.",
            "Homarus gammarus is a species of clawed lobster.",
        ),
    )
    assert chunks_for(document) == [
        ("field-guide-p0001", "The European lobster is Homarus gammarus."),
        ("field-guide-p0002", "Homarus gammarus is a species of clawed lobster."),
    ]


def test_an_outline_keeps_the_section_title_on_each_item():
    text = "\n".join([
        "page 1",
        "3-28. Travel.",
        "a. Maximum Advance. The advance is 500 dollars.",
        "b. Receipts. A receipt is required.",
        "3-29. Leave.",
        "a. Notice. Notice is required.",
        "b. Approval. A manager approves the leave.",
    ])
    chunks = structured_paragraphs(text)
    assert any(chunk.startswith("3-28. Travel.") and "Maximum Advance" in chunk and "Receipts" in chunk for chunk in chunks)
    assert all("page 1" not in chunk for chunk in chunks)


def test_plain_text_html_and_a_folder_keep_every_paragraph(tmp_path: Path):
    (tmp_path / "guide.txt").write_text(
        "The European lobster is Homarus gammarus.\n\nHomarus gammarus is a species of clawed lobster.\n",
        encoding="utf-8",
    )
    (tmp_path / "page.html").write_text(
        "<html><body><p>The crab is Cancer pagurus.</p><script>ignore()</script></body></html>",
        encoding="utf-8",
    )
    documents = {item.name: item.paragraphs for item in open_source(tmp_path)}
    assert documents["guide"] == (
        "The European lobster is Homarus gammarus.",
        "Homarus gammarus is a species of clawed lobster.",
    )
    assert documents["page"] == ("The crab is Cancer pagurus.",)
    assert paragraphs_from_text("The European lob-\nster is Homarus gammarus.") == [
        "The European lobster is Homarus gammarus."
    ]


def test_pdf_and_word_extract_the_written_paragraph(tmp_path: Path):
    pdf_path = tmp_path / "lobster.pdf"
    pdf_path.write_bytes(_pdf("The European lobster is Homarus gammarus."))
    assert "Homarus gammarus" in PdfReader(str(pdf_path)).pages[0].extract_text()
    found = open_source(pdf_path)
    assert found[0].paragraphs == ("The European lobster is Homarus gammarus.",)

    word = docx.Document()
    word.add_paragraph("The European lobster is Homarus gammarus.")
    word.add_paragraph("Homarus gammarus is a species of clawed lobster.")
    word_path = tmp_path / "lobster.docx"
    word.save(word_path)
    docx_doc = open_source(word_path)[0]
    assert docx_doc.paragraphs == (
        "The European lobster is Homarus gammarus.",
        "Homarus gammarus is a species of clawed lobster.",
    )


def test_a_sheet_stores_the_row_and_the_connection(tmp_path: Path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "links"
    sheet.append(["from", "relation", "to"])
    sheet.append(["lobster", "maps to", "Homarus gammarus"])
    path = tmp_path / "links.xlsx"
    book.save(path)
    book.close()
    document = open_source(path)[0]
    assert document.paragraphs == ("lobster maps to Homarus gammarus.",)
    assert document.connections == (("lobster", "Homarus gammarus", "maps to"),)

    csv_path = tmp_path / "notes.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["species", "habitat"])
        writer.writerow(["European lobster", "Atlantic"])
    table = open_source(csv_path)[0]
    assert table.connections == ()
    assert table.paragraphs == ("species: European lobster. habitat: Atlantic.",)


def test_an_epub_chapter_is_read_and_a_legacy_file_is_refused(tmp_path: Path):
    epub = tmp_path / "book.epub"
    with zipfile.ZipFile(epub, "w") as package:
        package.writestr(
            "chapter.xhtml",
            "<html><body><p>The European lobster is Homarus gammarus.</p></body></html>",
        )
    assert open_source(epub)[0].paragraphs == ("The European lobster is Homarus gammarus.",)
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"not a document")
    with pytest.raises(ValueError, match="docx"):
        open_source(legacy)


def test_sheet_connections_are_added_beside_existing_edges(tmp_path: Path):
    path = tmp_path / "word_relations.json"
    assert add_edges(path, [("lobster", "Homarus gammarus")]) == 1
    assert add_edges(path, [("lobster", "Homarus gammarus"), ("crab", "Cancer pagurus")]) == 1
    pairs = {(edge.from_word, edge.to_word) for edge in load_edges(path)}
    assert ("homarus gammarus", "lobster") in pairs
    assert ("cancer pagurus", "crab") in pairs


def _pdf(text: str) -> bytes:
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def test_a_shorter_source_name_does_not_share_its_prefix():
    from cpu_llm.documents import replace_prefix

    assert replace_prefix(SourceDocument("hr", ("x",))) == "hr-p"
    assert chunks_for(SourceDocument("hr-policy", ("y",)))[0][0] == "hr-policy-p0001"


def test_two_names_cannot_claim_one_slug(tmp_path):
    registry = tmp_path / "source_slugs.json"
    claim_source_names(registry, ["Q3 Report.pdf"])
    before = registry.read_bytes()
    with pytest.raises(ValueError, match="already used"):
        claim_source_names(registry, ["q3-report.pdf"])
    assert registry.read_bytes() == before
    claim_source_names(registry, ["Q3 Report.pdf"])
    claim_source_names(registry, ["q3 report.PDF"])
    owners = json.loads(registry.read_text(encoding="utf-8"))
    assert owners[_slug("Q3 Report.pdf")] == "q3 report.pdf"


def test_a_colliding_ingest_stops_before_any_paragraph_is_stored(tmp_path):
    from cpu_llm.pipeline import Appliance

    app = Appliance.__new__(Appliance)
    app.source_slug_path = tmp_path / "source_slugs.json"
    wrote: list[str] = []

    def catalog_chunks(chunks, propositions, replace_prefix=None):
        wrote.append(replace_prefix)
        return len(chunks)

    app.catalog_chunks = catalog_chunks
    app.ingest(SourceDocument("Q3 Report.pdf", ("Quarterly numbers.",)))
    assert wrote == ["q3-report-pdf-p"]
    before = app.source_slug_path.read_bytes()
    with pytest.raises(ValueError, match="already used"):
        app.ingest(SourceDocument("q3-report.pdf", ("Other numbers.",)))
    assert app.source_slug_path.read_bytes() == before
    assert wrote == ["q3-report-pdf-p"]
    app.ingest(SourceDocument("q3 report.PDF", ("Quarterly numbers again.",)))
    assert wrote == ["q3-report-pdf-p", "q3-report-pdf-p"]


def test_concurrent_claims_keep_the_registry(tmp_path):
    registry = tmp_path / "source_slugs.json"

    def claim_many(path: Path, names: list[str]) -> list[BaseException]:
        barrier = threading.Barrier(len(names))
        errors: list[BaseException] = []

        def claim(name: str) -> None:
            barrier.wait()
            try:
                claim_source_names(path, [name])
            except BaseException as exc:
                errors.append(exc)

        threads = [threading.Thread(target=claim, args=(name,)) for name in names]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        return errors

    distinct = [f"ledger-{number}.pdf" for number in range(8)]
    assert claim_many(registry, distinct) == []
    stored = json.loads(registry.read_text(encoding="utf-8"))
    assert set(stored) == {_slug(name) for name in distinct}

    claim_source_names(registry, ["Q3 Report.pdf"])
    for _ in range(20):
        errors = claim_many(registry, ["q3-report.pdf", "Q3-Report.PDF"])
        assert len(errors) == 2
        assert all(isinstance(item, ValueError) for item in errors)
        stored = json.loads(registry.read_text(encoding="utf-8"))
        assert stored[_slug("Q3 Report.pdf")] == "q3 report.pdf"
        assert len(stored) == len(distinct) + 1

    for round_number in range(20):
        fresh = tmp_path / f"race-{round_number}.json"
        errors = claim_many(fresh, ["Q3 Report.pdf", "q3-report.pdf"])
        assert len(errors) == 1
        assert isinstance(errors[0], ValueError)
        stored = json.loads(fresh.read_text(encoding="utf-8"))
        assert list(stored) == ["q3-report-pdf"]
        assert stored["q3-report-pdf"] in {"q3 report.pdf", "q3-report.pdf"}


def test_a_truncated_name_keeps_a_hash_of_the_full_name():
    left = "quarterly-operating-review-" + ("alpha" * 20)
    right = "quarterly-operating-review-" + ("beta" * 20)
    assert _slug(left) != _slug(right)
    assert len(_slug(left)) <= 72
    assert len(_slug(right)) <= 72


def test_a_name_that_looks_like_a_paragraph_id_is_rejected():
    with pytest.raises(ValueError, match="paragraph id"):
        chunks_for(SourceDocument("report-p0003", ("x",)))
