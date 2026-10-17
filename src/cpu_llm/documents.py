"""Turn a developer's source into the paragraphs ``catalog_chunks`` stores.

A source is either a ``SourceDocument`` the caller already split, or a file
the reader splits. Each paragraph becomes ``{name}-p0001``, ``{name}-p0002``,
and so on. Those ids are the chunks a later question retrieves.

Accepted files are ``.txt``, ``.md``, ``.html``, ``.pdf``, ``.docx``,
``.csv``, ``.xlsx``, and ``.epub``. A directory is every accepted file inside
it. A spreadsheet row whose headers name two ends becomes a paragraph and a
connection between those two cells.
"""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

_ACCEPTED = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".xlsm",
    ".epub",
}
_FROM = {"from", "source", "left", "parent", "subject"}
_TO = {"to", "target", "right", "child", "object"}
_RELATION = {"relation", "relationship", "connection", "maps to", "map", "predicate", "link"}


@dataclass(frozen=True)
class SourceDocument:
    """Paragraphs of one named source, plus any connections a sheet stated.

    ``connections`` are ``(left, right, label)`` taken from column headers.
    The label stays in the paragraph text. The stored word edge is the pair.
    """

    name: str
    paragraphs: tuple[str, ...]
    connections: tuple[tuple[str, str, str], ...] = field(default_factory=tuple)


def open_source(source: str | Path | SourceDocument | Iterable[SourceDocument | str | Path]) -> list[SourceDocument]:
    """Read one file, one directory, one document, or a list of those."""
    if isinstance(source, SourceDocument):
        return [_checked(source)]
    if isinstance(source, (str, Path)):
        return _open_path(Path(source))
    documents: list[SourceDocument] = []
    for item in source:
        documents.extend(open_source(item))
    if not documents:
        raise ValueError("the source has no paragraphs")
    return documents


def chunks_for(document: SourceDocument) -> list[tuple[str, str]]:
    """``(document_id, paragraph)`` pairs in paragraph order."""
    slug = _slug(document.name)
    return [
        (f"{slug}-p{number:04d}", paragraph)
        for number, paragraph in enumerate(document.paragraphs, start=1)
    ]


def replace_prefix(document: SourceDocument) -> str:
    """Prefix that replaces an earlier ingest of this same source."""
    return f"{_slug(document.name)}-p"


def paragraphs_from_text(text: str) -> list[str]:
    """Paragraphs split on a blank line, with wrapped lines joined back."""
    cleaned = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    paragraphs: list[str] = []
    for block in re.split(r"\n\s*\n", cleaned):
        lines = [re.sub(r"\s+", " ", line).strip() for line in block.split("\n")]
        lines = [line for line in lines if line]
        if not lines:
            continue
        current = lines[0]
        for line in lines[1:]:
            if current[-1:] in ".!?:;" and line[:1].isupper():
                paragraphs.append(current.strip())
                current = line
                continue
            if current.endswith("-"):
                current = current[:-1] + line
            else:
                current = f"{current} {line}"
        if current.strip():
            paragraphs.append(current.strip())
    return paragraphs


_FURNITURE = re.compile(
    r"^\s*(?:_{4,}|={4,}.*|\d{1,2}-\d{1,3}|\d{1,2}/\d{2,4}|[ivxlc]{1,5}|page \d+(?: of \d+)?|"
    r"\d{1,2}/\d{2,4}\s+\d{1,2}-\d{1,3}|\d{1,2}-\d{1,3}\s+\d{1,2}/\d{2,4})\s*$",
    re.IGNORECASE,
)
_LABEL = re.compile(r"^\s*(?P<label>\d{1,2}[-.]\d{1,3}\.?|[a-z]\.|\(\d{1,2}\)|\([a-z]\))\s+(?P<rest>\S.*)$")
_CHAPTER = re.compile(r"^\s*(?:CHAPTER|SECTION)\s+\w+\.?\s+.+$")
_MAX_CHUNK = 1400


def _label_level(label: str) -> int:
    if re.match(r"\d{1,2}[-.]\d{1,3}\.?$", label):
        return 1
    if re.match(r"[a-z]\.$", label):
        return 2
    if re.match(r"\(\d{1,2}\)$", label):
        return 3
    return 4


def structured_paragraphs(text: str) -> list[str]:
    """Chunks for an outline extracted from a PDF.

    Page numbers, dates, and rules are dropped. Hard-wrapped lines are joined.
    Numbered items become units, and a list header ending in ':' stays with its
    items. Every nested chunk starts with its section title. A document without
    that numbering falls back to ordinary paragraphs.
    """
    lines = [ln.rstrip() for ln in text.replace("\x00", " ").replace("\r", "\n").split("\n")]
    lines = [ln for ln in lines if ln.strip() and not _FURNITURE.match(ln)]
    items: list[dict] = []
    for raw in lines:
        line = re.sub(r"\s+", " ", raw).strip()
        found = _LABEL.match(line)
        if found:
            items.append({"label": found["label"], "level": _label_level(found["label"]), "text": found["rest"]})
        elif _CHAPTER.match(line):
            items.append({"label": "", "level": 0, "text": line})
        elif items:
            prev = items[-1]
            prev["text"] = prev["text"][:-1] + line if prev["text"].endswith("-") else f"{prev['text']} {line}"
        else:
            items.append({"label": "", "level": 0, "text": line})
    if sum(1 for it in items if it["level"] in {1, 2}) < 5:
        return paragraphs_from_text(text)

    chunks: list[str] = []
    section = ""
    index = 0
    while index < len(items):
        item = items[index]
        if item["level"] == 0:
            index += 1
            continue
        if item["level"] == 1:
            title = re.split(r"(?<=[.])\s", item["text"], maxsplit=1)[0].rstrip(".")
            section = f"{item['label'].rstrip('.')} {title}".strip()
        body = f"{item['label']} {item['text']}".strip()
        nxt = index + 1
        if item["text"].rstrip().endswith(":") or len(item["text"]) < 60:
            while nxt < len(items) and items[nxt]["level"] > item["level"]:
                body += f" {items[nxt]['label']} {items[nxt]['text']}"
                nxt += 1
        prefix = f"{section}. " if section and item["level"] > 1 else ""
        text_chunk = (prefix + body).strip()
        while len(text_chunk) > _MAX_CHUNK:
            cut = text_chunk.rfind(". ", 0, _MAX_CHUNK) + 1 or _MAX_CHUNK
            chunks.append(text_chunk[:cut].strip())
            text_chunk = (prefix + text_chunk[cut:].strip()).strip()
        chunks.append(text_chunk)
        index = nxt if nxt > index + 1 else index + 1
    return [chunk for chunk in chunks if len(chunk) > 15]


def _open_path(path: Path) -> list[SourceDocument]:
    if not path.exists():
        raise FileNotFoundError(path)
    if path.is_dir():
        files = [item for item in sorted(path.rglob("*")) if item.is_file() and item.suffix.lower() in _ACCEPTED]
        if not files:
            raise ValueError(f"{path} has no accepted file ({', '.join(sorted(_ACCEPTED))})")
        documents: list[SourceDocument] = []
        for item in files:
            documents.extend(_open_file(item))
        return documents
    return _open_file(path)


def _open_file(path: Path) -> list[SourceDocument]:
    suffix = path.suffix.lower()
    if suffix in {".doc", ".xls", ".ppt"}:
        raise ValueError(f"{path.name} is a legacy Office file; save it as {suffix}x")
    if suffix not in _ACCEPTED:
        accepted = ", ".join(sorted(_ACCEPTED))
        raise ValueError(f"{path.name} is not an accepted source. Accepted files: {accepted}")
    if suffix in {".txt", ".md", ".markdown"}:
        documents = [_text_document(path.stem, path.read_text(encoding="utf-8"))]
    elif suffix in {".html", ".htm"}:
        documents = [_text_document(path.stem, _html_text(path.read_text(encoding="utf-8")))]
    elif suffix == ".pdf":
        documents = [SourceDocument(path.stem, tuple(structured_paragraphs(_pdf_text(path))))]
    elif suffix == ".docx":
        documents = [_text_document(path.stem, _docx_text(path))]
    elif suffix == ".epub":
        documents = [_text_document(path.stem, _epub_text(path))]
    elif suffix == ".csv":
        documents = [_table_document(path.stem, _csv_rows(path))]
    else:
        documents = _xlsx_documents(path)
    return [_checked(document) for document in documents]


def _text_document(name: str, text: str) -> SourceDocument:
    return SourceDocument(name, tuple(paragraphs_from_text(text)))


def _checked(document: SourceDocument) -> SourceDocument:
    paragraphs = tuple(paragraph.strip() for paragraph in document.paragraphs if paragraph and paragraph.strip())
    if not document.name.strip():
        raise ValueError("a source needs a name")
    if not paragraphs:
        raise ValueError(f"{document.name} has no paragraphs")
    return SourceDocument(document.name.strip(), paragraphs, document.connections)


def _slug(name: str) -> str:
    path = Path(name)
    suffix = path.suffix.lower().lstrip(".")
    stem = path.stem if suffix else name
    raw = f"{stem}-{suffix}" if suffix else stem
    slug = re.sub(r"[^a-z0-9]+", "-", raw.casefold()).strip("-")
    return (slug or "source")[:72]


def _html_text(html: str) -> str:
    parser = _HTMLText()
    parser.feed(html)
    return "".join(parser.parts)


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip += 1
        if tag in {"p", "div", "h1", "h2", "h3", "h4", "li", "br", "tr"}:
            self.parts.append("\n\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip:
            self._skip -= 1
        if tag in {"p", "h1", "h2", "h3", "h4", "li", "tr"}:
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _epub_text(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        names = [name for name in package.namelist() if name.lower().endswith((".xhtml", ".html", ".htm"))]
        if not names:
            raise ValueError(f"{path.name} has no HTML chapter")
        chapters = []
        for name in sorted(names):
            html = package.read(name).decode("utf-8", errors="replace")
            chapters.append(_html_text(html))
    return "\n\n".join(chapters)


def _csv_rows(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return [list(row) for row in csv.reader(handle) if any(cell.strip() for cell in row)]


def _xlsx_documents(path: Path) -> list[SourceDocument]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    documents: list[SourceDocument] = []
    try:
        for sheet in book.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None and str(cell).strip() for cell in row)
            ]
            if rows:
                documents.append(_table_document(f"{path.stem}-{sheet.title}", rows))
    finally:
        book.close()
    return documents


def _table_document(name: str, rows: list[list[str]]) -> SourceDocument:
    if not rows:
        return SourceDocument(name, ())
    header = [cell.strip() for cell in rows[0]]
    keyed = _header_map(header) if len(rows) > 1 else None
    paragraphs: list[str] = []
    connections: list[tuple[str, str, str]] = []
    body = rows[1:] if len(rows) > 1 else rows
    labels = header if len(rows) > 1 else []
    for row in body:
        cells = [cell.strip() for cell in row]
        if not any(cells):
            continue
        if keyed:
            left = _cell(cells, keyed.get("from"))
            right = _cell(cells, keyed.get("to"))
            label = _cell(cells, keyed.get("relation")) or "connects to"
            if left and right:
                connections.append((left, right, label))
                paragraphs.append(f"{left} {label} {right}.")
                continue
        if labels:
            pairs = [f"{label}: {value}" for label, value in zip(labels, cells) if label and value]
            if pairs:
                paragraphs.append(". ".join(pairs) + ".")
        else:
            text = ". ".join(cell for cell in cells if cell)
            if text:
                paragraphs.append(text if text.endswith(".") else f"{text}.")
    return SourceDocument(name, tuple(paragraphs), tuple(connections))


def _header_map(header: list[str]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, cell in enumerate(header):
        name = cell.casefold()
        if name in _FROM and "from" not in found:
            found["from"] = index
        elif name in _TO and "to" not in found:
            found["to"] = index
        elif name in _RELATION and "relation" not in found:
            found["relation"] = index
    if "from" in found and "to" in found:
        return found
    return None


def _cell(cells: list[str], index: int | None) -> str:
    if index is None or index >= len(cells):
        return ""
    return cells[index].strip()
